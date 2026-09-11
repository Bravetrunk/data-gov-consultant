"""
Word Document Generator for data-gov-consultant Platform.
Generates comprehensive, publication-grade Transformation Master Blueprint (.docx)
with executive headers, formal tables, and ready-to-sign appointment charters.
"""

import os
from typing import List, Dict, Any
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

from models.state import ClientOrganization, UseCaseItem, RACIItem, MetadataRecord, CloudTCOComparison


def set_cell_background(cell, fill_hex):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)


def build_master_transformation_report(
    output_path: str,
    client: ClientOrganization,
    maturity_score: int,
    maturity_narrative: str,
    use_cases: List[UseCaseItem],
    raci_items: List[RACIItem],
    tco_list: List[CloudTCOComparison]
) -> str:
    doc = Document()

    # Configure Normal Style font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Aptos'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x33, 0x33, 0x33)

    # 1. Cover / Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(36)
    p_title.paragraph_format.space_after = Pt(12)
    run_title = p_title.add_run(f"พิมพ์เขียวการเปลี่ยนผ่านสู่องค์กรขับเคลื่อนด้วยข้อมูลและเอไอ\n(Enterprise Data & AI Transformation Master Blueprint)")
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(24)
    run_sub = p_sub.add_run(f"จัดทำสำหรับ: {client.name}\n(อุตสาหกรรม: {client.industry.upper()} | ขนาด: {client.organization_size})")
    run_sub.font.size = Pt(14)
    run_sub.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(48)
    run_meta = p_meta.add_run("คณะผู้จัดทำ: data-gov-consultant AI-Native Transformation Practice\nมาตรฐานอ้างอิง: สพร. (DGA), DAMA-DMBOK, PDPA 2562, ISO/IEC 42001, NIST AI RMF 1.0")
    run_meta.font.size = Pt(10)
    run_meta.font.italic = True

    doc.add_page_break()

    # 2. Executive Summary
    h1 = doc.add_heading("1. บทสรุปสำหรับผู้บริหาร (Executive Summary)", level=1)
    h1.paragraph_format.space_before = Pt(18)
    
    p = doc.add_paragraph(
        f"รายงานฉบับนี้จัดทำขึ้นเพื่อเป็นแผนแม่บทเชิงกลยุทธ์และสถาปัตยกรรมทางเทคนิคในการยกระดับ {client.name} "
        f"จากการบริหารจัดการข้อมูลแบบเดิม สู่การมีธรรมาภิบาลข้อมูล (Data Governance) ที่มั่นคงปลอดภัย ถูกต้องตามกฎหมาย "
        f"และต่อยอดสู่สถาปัตยกรรม Big Data Lakehouse รวมถึงการประยุกต์ใช้โมเดลปัญญาประดิษฐ์ (AI Transformation) "
        f"อย่างมีจริยธรรมและมีประสิทธิภาพสูงสุด โดยยึดหลัก Business-First และการสร้างผลตอบแทนจากการลงทุน (ROI) เป็นสำคัญ"
    )
    p.paragraph_format.line_spacing = 1.25

    # 3. As-Is Assessment & Maturity
    h2 = doc.add_heading("2. การประเมินสถานะปัจจุบันและระดับวุฒิภาวะ (AS-IS Maturity Assessment)", level=1)
    doc.add_paragraph(
        f"จากการประเมินตามกรอบแนวคิดธรรมาภิบาลข้อมูลภาครัฐของ สพร. (DGA Readiness Assessment 6 ระดับ) "
        f"พบว่าปัจจุบัน {client.name} มีระดับความพร้อมอยู่ที่:"
    )
    
    p_score = doc.add_paragraph()
    p_score.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_score = p_score.add_run(f"ระดับที่ {maturity_score} : {maturity_narrative}")
    r_score.font.size = Pt(16)
    r_score.font.bold = True
    r_score.font.color.rgb = RGBColor(0x1F, 0x49, 0x7D)

    doc.add_paragraph("ปัญหาและอุปสรรคสำคัญที่พบในปัจจุบัน (Key Pain Points):")
    for pt in client.primary_pain_points:
        p_pt = doc.add_paragraph(pt, style='List Bullet')
        p_pt.paragraph_format.space_after = Pt(4)

    # 4. Use-Case Portfolio & Financial ROI
    h3 = doc.add_heading("3. พอร์ตโฟลิโอยูสเคสธุรกิจและประมาณการความคุ้มค่า (Use-Case Portfolio & ROI)", level=1)
    doc.add_paragraph(
        "การจัดลำดับความสำคัญของ Use Cases ทางธุรกิจด้วยเมทริกซ์ 2x2 (Business Impact vs. Technical Feasibility) "
        "เพื่อเฟ้นหาโครงการ Quick Win ที่สามารถสร้างผลกำไรหรือลดต้นทุนได้ทันทีภายใน 6–12 เดือนแรก:"
    )

    # Table of Use Cases
    table_uc = doc.add_table(rows=1, cols=6)
    table_uc.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers_uc = ["รหัส", "ชื่อโครงการ Use Case", "หน่วยงาน", "ผลประโยชน์/ปี (ลบ.)", "ต้นทุน (ลบ.)", "คืนทุน (เดือน)"]
    for idx, name in enumerate(headers_uc):
        cell = table_uc.rows[0].cells[idx]
        cell.text = name
        set_cell_background(cell, "1F497D")
        p_hdr = cell.paragraphs[0]
        p_hdr.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_hdr.runs[0].font.bold = True
        p_hdr.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for uc in use_cases:
        row_cells = table_uc.add_row().cells
        row_cells[0].text = uc.id
        row_cells[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        row_cells[1].text = uc.title
        row_cells[2].text = uc.business_unit
        row_cells[3].text = f"{uc.estimated_annual_benefit_thb / 1_000_000:.2f}"
        row_cells[3].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row_cells[4].text = f"{uc.implementation_cost_thb / 1_000_000:.2f}"
        row_cells[4].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row_cells[5].text = f"{uc.payback_months:.1f}"
        row_cells[5].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        for c in row_cells:
            set_cell_margins(c)

    # 5. Data Governance Framework & 6-Stage Lifecycle
    h4 = doc.add_heading("4. กรอบธรรมาภิบาลข้อมูลและโครงสร้างบทบาท (Governance & Operating Model)", level=1)
    doc.add_paragraph(
        "การขับเคลื่อนธรรมาภิบาลข้อมูลตลอดวงจรชีวิตข้อมูล 6 ขั้นตอน (Data Lifecycle Management) "
        "พร้อมโครงสร้างคณะกรรมการและคณะทำงานที่ชัดเจน:"
    )

    lifecycle_steps = [
        ("1. การสร้างข้อมูล (Create):", "กำหนดให้สร้างข้อมูลจากแหล่งที่เชื่อถือได้ ห้ามบันทึกข้อมูลเท็จตาม พ.ร.บ. คอมพิวเตอร์ และทำ Metadata ตั้งแต่ต้นทาง"),
        ("2. การจัดเก็บข้อมูล (Store):", "จัดเก็บในระบบฐานข้อมูลที่มีการเข้ารหัส (Encryption at Rest) แยกฟิลด์ความลับ และปฏิบัติตามระยะเวลาจัดเก็บ (Retention Policy)"),
        ("3. การใช้/ประมวลผล (Use):", "จำกัดสิทธิ์ตามหน้าที่ (Role-Based Access Control) ห้ามนำข้อมูลลูกค้า/ผู้ป่วยไปใช้ในพื้นที่สาธารณะหรือแสวงหาผลประโยชน์ส่วนตัว"),
        ("4. การเปิดเผยข้อมูล (Publish):", "คัดกรองข้อมูลสาธารณะ (Open Data) ขึ้นสู่ Data Catalog โดยต้องไม่มีข้อมูลส่วนบุคคลอ่อนไหว (PDPA ม.26) ปรากฏอยู่"),
        ("5. การจัดเก็บถาวร (Archive):", "คัดลอกข้อมูลที่หมดช่วงใช้งานเข้าสู่ Cold Storage พร้อมแผนทดสอบการกู้คืน (Restore Drill) อย่างน้อยปีละ 1 ครั้ง"),
        ("6. การทำลายข้อมูล (Destroy):", "ตั้งคณะกรรมการอนุมัติทำลายเมื่อพ้นกำหนดระยะเวลาตามกฎหมาย พร้อมเก็บบันทึกหลักฐาน (Audit Log) ไว้อย่างน้อย 1 ปี")
    ]
    for step_title, step_desc in lifecycle_steps:
        p_step = doc.add_paragraph()
        r_step = p_step.add_run(step_title + " ")
        r_step.font.bold = True
        p_step.add_run(step_desc)

    # 6. Big Data Architecture & Cloud TCO
    h5 = doc.add_heading("5. สถาปัตยกรรมบิ๊กดาต้าและการเปรียบเทียบต้นทุนคลาวด์ (Architecture & Cloud TCO)", level=1)
    doc.add_paragraph(
        "พิมพ์เขียว Modern Data Lakehouse สำหรับรองรับทั้งงานรายงานอัจฉริยะ (BI) และโมเดล Machine Learning "
        "โดยสรุปเปรียบเทียบต้นทุนรวมในการเป็นเจ้าของ (3-Year TCO) ระหว่างทางเลือกต่าง ๆ:"
    )

    table_tco = doc.add_table(rows=1, cols=4)
    table_tco.alignment = WD_TABLE_ALIGNMENT.CENTER
    tco_headers = ["สถาปัตยกรรมคลาวด์", "ค่าบริการรายปี (ลบ.)", "TCO สะสม 3 ปี (ลบ.)", "ข้อได้เปรียบหลัก"]
    for idx, name in enumerate(tco_headers):
        cell = table_tco.rows[0].cells[idx]
        cell.text = name
        set_cell_background(cell, "2C3E50")
        p_h = cell.paragraphs[0]
        p_h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_h.runs[0].font.bold = True
        p_h.runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    for t in tco_list:
        row_cells = table_tco.add_row().cells
        row_cells[0].text = t.provider
        row_cells[1].text = f"{t.annual_total_thb / 1_000_000:.2f}"
        row_cells[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        row_cells[2].text = f"{t.three_year_tco_thb / 1_000_000:.2f}"
        row_cells[2].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        if row_cells[2].paragraphs[0].runs:
            row_cells[2].paragraphs[0].runs[0].font.bold = True
        row_cells[3].text = t.pros
        for c in row_cells:
            set_cell_margins(c)

    # 7. AI Transformation & ISO 42001
    h6 = doc.add_heading("6. ธรรมาภิบาลปัญญาประดิษฐ์ (AI Governance & ISO/IEC 42001)", level=1)
    doc.add_paragraph(
        "แนวทางการกำกับดูแลโมเดล AI และ Generative AI เพื่อป้องกันความเสี่ยงด้านภาพลักษณ์และกฎหมาย:"
    )
    ai_pillars = [
        ("Fairness & Non-Discrimination:", "ทดสอบโมเดลว่าไม่มีอคติต่อเพศ อายุ หรือถิ่นที่อยู่ของผู้รับบริการ"),
        ("Explainability (XAI):", "โมเดลที่กระทบต่อการตัดสินใจสำคัญ ต้องสามารถอธิบายเหตุผลของผลลัพธ์ได้"),
        ("Human-in-the-Loop:", "กำหนดให้ผู้เชี่ยวชาญ (แพทย์/หัวหน้างาน) เป็นผู้อนุมัติขั้นสุดท้ายในเคสสำคัญ"),
        ("Training Data Governance:", "ข้อมูลที่นำมาเทรนต้องผ่านการทำ Anonymization ตัดตัวตน 100% และตรวจสอบลิขสิทธิ์"),
        ("Prompt Injection & Leakage Defense:", "วางเกราะป้องกันไม่ให้พนักงานป้อนข้อมูลความลับองค์กรลงในระบบ Public AI")
    ]
    for p_name, p_detail in ai_pillars:
        p_ai = doc.add_paragraph()
        r_ai = p_ai.add_run(p_name + " ")
        r_ai.font.bold = True
        p_ai.add_run(p_detail)

    # Annex: Ready-to-Sign Appointment Template
    doc.add_page_break()
    h_app = doc.add_heading("ภาคผนวก: ร่างคำสั่งแต่งตั้งคณะกรรมการธรรมาภิบาลข้อมูล (พร้อมลงนาม)", level=1)
    
    p_order = doc.add_paragraph()
    p_order.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_ord = p_order.add_run(f"คำสั่ง {client.name}\nที่ ..... / 2569\nเรื่อง แต่งตั้งคณะกรรมการธรรมาภิบาลข้อมูล (Data Governance Council)")
    r_ord.font.bold = True
    r_ord.font.size = Pt(13)

    doc.add_paragraph(
        f"เพื่อให้การบริหารจัดการข้อมูลและการขับเคลื่อนนวัตกรรมปัญญาประดิษฐ์ของ {client.name} "
        "มีความมั่นคงปลอดภัย มีคุณภาพ และสอดคล้องตามพระราชบัญญัติคุ้มครองข้อมูลส่วนบุคคล พ.ศ. 2562 "
        "จึงมีคำสั่งแต่งตั้งคณะกรรมการธรรมาภิบาลข้อมูล ประกอบด้วย:"
    )
    
    council_roles = [
        ("1. ประธานเจ้าหน้าที่บริหาร (CEO) / ผู้อำนวยการ", "ประธานกรรมการ"),
        ("2. ผู้บริหารเทคโนโลยีสารสนเทศระดับสูง (CIO/CTO)", "รองประธานกรรมการ"),
        ("3. เจ้าหน้าที่คุ้มครองข้อมูลส่วนบุคคล (DPO / ฝ่ายกฎหมาย)", "กรรมการ"),
        ("4. ผู้อำนวยการสายงานธุรกิจและปฏิบัติการ", "กรรมการ"),
        ("5. หัวหน้าคณะบริกรข้อมูล (Lead Data Steward)", "กรรมการและเลขานุการ")
    ]
    for role_name, pos in council_roles:
        p_c = doc.add_paragraph()
        p_c.add_run(f"• {role_name} ").font.bold = True
        p_c.add_run(f"ดำรงตำแหน่ง {pos}")

    p_sign = doc.add_paragraph()
    p_sign.paragraph_format.space_before = Pt(36)
    p_sign.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sign.add_run("สั่ง ณ วันที่ ........................................\n\n\nลงนาม .....................................................\n( ..................................................... )\nตำแหน่ง ผู้มีอำนาจลงนาม")

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
